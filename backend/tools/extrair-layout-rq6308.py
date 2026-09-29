"""Extrai o layout fiel da ficha RQ 6308 (.xlsm oficial) para o frontend.

Uso: python backend/tools/extrair-layout-rq6308.py <arquivo.xlsm> <saida.js>

Gera um JS que define window.RQ6308_LAYOUT com larguras, alturas, merges,
estilos resolvidos (cores de tema já convertidas para hex), fórmulas,
formatos numéricos, células desbloqueadas, logo (data URI) e posição do gráfico.
"""
import base64
import colorsys
import json
import sys
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter, range_boundaries

EMU_POR_PX = 9525
ULTIMA_COLUNA = 17  # Q
ULTIMA_LINHA = 53   # área de impressão A1:Q53

NS_A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def cores_do_tema(caminho):
    xml = zipfile.ZipFile(caminho).read('xl/theme/theme1.xml')
    raiz = ET.fromstring(xml)
    esquema = raiz.find(f'.//{NS_A}clrScheme')
    cores = {}
    for filho in esquema:
        nome = filho.tag.replace(NS_A, '')
        srgb = filho.find(f'{NS_A}srgbClr')
        sys_clr = filho.find(f'{NS_A}sysClr')
        cores[nome] = (srgb.get('val') if srgb is not None else sys_clr.get('lastClr'))
    # Ordem de índice de tema do Excel: lt1, dk1, lt2, dk2, accent1..6, hlink, folHlink
    ordem = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
    return [cores[n] for n in ordem]


def aplicar_tint(hex6, tint):
    if not tint:
        return hex6
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return ''.join(f'{round(c * 255):02X}' for c in (r, g, b))


def resolver_cor(cor, tema):
    if cor is None:
        return None
    tipo = cor.type
    if tipo == 'rgb' and isinstance(cor.rgb, str):
        return '#' + cor.rgb[-6:]
    if tipo == 'theme':
        return '#' + aplicar_tint(tema[cor.theme], cor.tint)
    if tipo == 'indexed':
        if cor.indexed in (64, 65):  # system foreground/background
            return None
        return '#' + COLOR_INDEX[cor.indexed][-6:]
    return None


LARGURA_BORDA = {
    'hair': '1px dotted', 'thin': '1px solid', 'dotted': '1px dotted', 'dashed': '1px dashed',
    'dashDot': '1px dashed', 'dashDotDot': '1px dashed', 'medium': '2px solid',
    'mediumDashed': '2px dashed', 'mediumDashDot': '2px dashed', 'mediumDashDotDot': '2px dashed',
    'slantDashDot': '2px dashed', 'thick': '3px solid', 'double': '3px double',
}


def borda(lado, tema):
    if lado is None or lado.style is None:
        return None
    cor = resolver_cor(lado.color, tema) or '#000000'
    return f'{LARGURA_BORDA.get(lado.style, "1px solid")} {cor}'


def estilo_celula(c, tema):
    s = {}
    f = c.font
    if f:
        s['fontFamily'] = f.name
        s['fontSize'] = f.sz
        if f.b:
            s['bold'] = True
        if f.i:
            s['italic'] = True
        if f.u:
            s['underline'] = True
        cor = resolver_cor(f.color, tema)
        if cor and cor.upper() != '#000000':
            s['color'] = cor
    if c.fill and c.fill.fill_type == 'solid':
        cor = resolver_cor(c.fill.fgColor, tema)
        if cor:
            s['fill'] = cor
    b = c.border
    for nome in ('left', 'right', 'top', 'bottom'):
        valor = borda(getattr(b, nome), tema)
        if valor:
            s[f'border_{nome}'] = valor
    a = c.alignment
    if a.horizontal:
        s['h'] = a.horizontal
    if a.vertical:
        s['v'] = a.vertical
    if a.wrap_text:
        s['wrap'] = True
    if a.shrink_to_fit:
        s['shrink'] = True
    if a.indent:
        s['indent'] = a.indent
    return s


def largura_px(chars):
    return int(((256 * chars + int(128 / 7)) / 256) * 7)


def main(origem, destino):
    tema = cores_do_tema(origem)
    wb = openpyxl.load_workbook(origem, keep_vba=True)
    ws = wb.active

    colunas = []
    for idx in range(1, ULTIMA_COLUNA + 1):
        letra = get_column_letter(idx)
        largura = ws.sheet_format.defaultColWidth or 8.43
        for dim in ws.column_dimensions.values():
            if dim.min and dim.max and dim.min <= idx <= dim.max and dim.width:
                largura = dim.width
        colunas.append({'letra': letra, 'px': largura_px(largura)})

    linhas = []
    for r in range(1, ULTIMA_LINHA + 1):
        dim = ws.row_dimensions.get(r)
        altura_pt = (dim.height if dim and dim.height else 12.75)
        linhas.append({'r': r, 'px': round(altura_pt * 4 / 3), 'oculta': bool(dim and dim.hidden)})

    merges = []
    for rng in ws.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(rng))
        if r1 > ULTIMA_LINHA or c1 > ULTIMA_COLUNA:
            continue
        merges.append({'ref': str(rng), 'c1': c1, 'r1': r1, 'c2': min(c2, ULTIMA_COLUNA), 'r2': min(r2, ULTIMA_LINHA)})

    celulas = {}
    for row in ws.iter_rows(min_row=1, max_row=ULTIMA_LINHA, min_col=1, max_col=ULTIMA_COLUNA):
        for c in row:
            if c.__class__.__name__ == 'MergedCell':
                # Merged filhas: só interessam as bordas (usadas nas extremidades do merge)
                s = {k: v for k, v in estilo_celula(c, tema).items() if k.startswith('border_')}
                if s:
                    celulas[c.coordinate] = {'s': s}
                continue
            item = {'s': estilo_celula(c, tema)}
            valor = c.value
            if isinstance(valor, str) and valor.startswith('='):
                item['f'] = valor
            elif valor is not None:
                item['v'] = valor
            if c.number_format and c.number_format != 'General':
                item['fmt'] = c.number_format
            celulas[c.coordinate] = item

    # O arquivo em uso na produção (foto de referência enviada pelo usuário em
    # 2026-09-29) tem todas as colunas de verificação D..P em pêssego; o modelo
    # .xlsm em branco alterna branco/pêssego nessas linhas.
    for r in [*range(15, 25), 28, *range(31, 41)]:
        for col in 'DEFGHIJKLMNOP':
            celulas.setdefault(f'{col}{r}', {'s': {}})['s']['fill'] = '#FCD5B5'

    img = ws._images[0]
    logo_bytes = zipfile.ZipFile(origem).read('xl/media/image1.png')
    de = img.anchor._from
    ate = img.anchor.to
    logo = {
        'dataUri': 'data:image/png;base64,' + base64.b64encode(logo_bytes).decode('ascii'),
        'de': {'col': de.col, 'colOffPx': round(de.colOff / EMU_POR_PX), 'row': de.row, 'rowOffPx': round(de.rowOff / EMU_POR_PX)},
        'ate': {'col': ate.col, 'colOffPx': round(ate.colOff / EMU_POR_PX), 'row': ate.row, 'rowOffPx': round(ate.rowOff / EMU_POR_PX)},
    }

    ch = ws._charts[0]
    grafico = {
        'de': {'col': ch.anchor._from.col, 'colOffPx': round(ch.anchor._from.colOff / EMU_POR_PX), 'row': ch.anchor._from.row, 'rowOffPx': round(ch.anchor._from.rowOff / EMU_POR_PX)},
        'ate': {'col': ch.anchor.to.col, 'colOffPx': round(ch.anchor.to.colOff / EMU_POR_PX), 'row': ch.anchor.to.row, 'rowOffPx': round(ch.anchor.to.rowOff / EMU_POR_PX)},
        'serie': 'D28:P28',
    }

    layout = {
        'aba': ws.title,
        'corSerieGrafico': '#' + tema[4],  # accent1: cor da 1ª série no estilo padrão do Excel
        'colunas': colunas,
        'linhas': linhas,
        'merges': merges,
        'celulas': celulas,
        'logo': logo,
        'grafico': grafico,
    }
    with open(destino, 'w', encoding='utf-8') as saida:
        saida.write('// Gerado por backend/tools/extrair-layout-rq6308.py a partir do .xlsm oficial. Não editar à mão.\n')
        saida.write('window.RQ6308_LAYOUT = ')
        json.dump(layout, saida, ensure_ascii=False, separators=(',', ':'))
        saida.write(';\n')
    print(f'ok: {len(celulas)} células, {len(merges)} merges -> {destino}')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
